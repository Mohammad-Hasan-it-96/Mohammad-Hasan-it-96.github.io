# -*- coding: utf-8 -*-
"""Build role-targeted CVs as PDF (Chrome) and DOCX (python-docx).

Usage:  python tools/build-cv.py [engineer|product] [en|ar]
        python tools/build-cv.py               # all four
Needs:  node (to read data.js for contact details), Chrome, python-docx.

Two variants, because one CV cannot serve both tracks:
  engineer — Laravel / React / TypeScript full-stack roles
  product  — product manager / product owner roles
Contact details come from data.js so they can never drift from the site."""
import io, json, os, subprocess, sys, tempfile, html

SP   = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SP)
OUT  = os.path.join(ROOT, "assets", "files")

def read_profile():
    js = ("eval(require('fs').readFileSync(process.argv[1],'utf8')"
          " + ';globalThis.__d={PROFILE};');"
          "process.stdout.write(JSON.stringify(globalThis.__d));")
    src = os.path.join(ROOT, "assets", "js", "data.js")
    r = subprocess.run(["node", "-e", js, src], capture_output=True)
    if r.returncode:
        sys.exit("node failed reading data.js:\n" + r.stderr.decode("utf-8", "replace"))
    return json.loads(r.stdout.decode("utf-8"))["PROFILE"]

PROF = read_profile()
LINKEDIN = "linkedin.com/in/mohamad-hasan-b78308274"
SITE     = "mohammad-hasan-it-96.github.io"
GITHUB   = "github.com/Mohammad-Hasan-it-96"

# ============================================================== CV CONTENT ===

CV = {}

# --------------------------------------------------- engineer / english ----
CV[("engineer", "en")] = {
    "title": "Full-Stack Developer — Laravel · React · TypeScript",
    "labels": {"summary": "Professional summary", "stack": "Core stack",
               "exp": "Professional experience", "proj": "Selected platforms",
               "edu": "Education", "lang": "Languages", "present": "Present"},
    "summary":
        "Full-stack developer with six years building commercial software that businesses run on daily. "
        "Laravel and PHP 8 on the back end — domain-driven service layers, versioned REST APIs, policy-based "
        "authorization, queues and caching — with React, Next.js and TypeScript on the front. I have taken four "
        "platforms and five mobile apps from first requirement to production support, including a multi-vendor "
        "marketplace and a licensing and billing control plane built as a 14-module modular monolith. "
        "I work in two full-stack tracks: Laravel with React, in production today, and full-stack TypeScript "
        "with NestJS, which I am building out now on the same architecture. "
        "Computer Systems Engineer, graduated with honours. Arabic-first (RTL) by default.",
    "stack": [
        ("Backend — Laravel", "PHP 8.4, Laravel 12, REST API design, Domain-Driven Design, modular monolith, service layer, "
                    "policies & RBAC, Sanctum / Passport, queues & scheduled jobs, Octane & Horizon, WebSockets, RabbitMQ"),
        ("Frontend — React & TypeScript", "TypeScript, React 18/19, Next.js 16, TanStack Query, Tailwind v4, "
                     "shadcn/ui, next-intl (RTL), Blade, JavaScript, HTML5, CSS3"),
        ("Full-stack JS (in progress)", "NestJS, Node.js, modules & dependency injection, DTOs & validation "
                     "pipes, guards & interceptors, TypeORM / Prisma, Jest — the same layered architecture, DTO "
                     "boundaries and role-based access I already run in Laravel, moved to the Node side. The React "
                     "and TypeScript front end is already in production."),
        ("Data & infrastructure", "MySQL, SQLite, Redis, Elasticsearch, schema design, query optimisation, "
                                  "Nginx, Apache, Linux, PM2, AWS S3, Cloudinary, Git"),
        ("Practices", "Layered architecture, DTO boundaries, domain events, API versioning, append-only migrations, "
                      "Larastan (max level), Pint, CI, code review"),
        ("Mobile", "Flutter, Dart, BLoC, Provider, get_it, Clean Architecture, Drift, sqflite, offline-first architecture"),
        ("Also", "Arduino, ESP32, C/C++, embedded control systems"),
    ],
    "exp": [
        {"r": "Full-Stack Developer (Independent)", "c": "Self-employed", "l": "Damascus, Syria · Remote",
         "d": "2021 – Present",
         "p": ["Own commercial products end to end for pharmacies, restaurants, retail shops and distributors — "
               "requirements, architecture, implementation, release and production support.",
               "Built EvoTech Core, a licensing, subscription and billing control plane: a modular monolith of 14 "
               "auto-discovered domain modules communicating only through contracts and domain events, with strict "
               "four-layer boundaries, DTO contracts, three authentication audiences and UUIDv7 route keys.",
               "Paired it with a Next.js 16 / React 19 / TypeScript front end — strict typing, TanStack Query, "
               "Tailwind v4 and next-intl, Arabic-RTL by default, deployed across separate subdomains via PM2 and Nginx.",
               "Enforced quality with Larastan at max level, Pint and CI; documented the API with Scramble (OpenAPI 3.1).",
               "Delivered five Flutter applications with offline-first SQLite/Drift data layers, versioned migrations, "
               "Bluetooth thermal printing and device-bound subscription licensing."]},
        {"r": "Full-Stack Laravel Developer", "c": "Rammaz Software", "l": "Remote — Turkey",
         "d": "2021 – May 2026",
         "p": ["Built and maintained scalable Laravel backends and versioned REST APIs serving web and mobile clients, "
               "in a fully remote engineering team.",
               "Shipped Trydos Market, a large multi-vendor e-commerce platform built with Domain-Driven Design — "
               "separate bounded contexts for Orders, Products, Wallet, Returns, Seller and Invoice — with POS, "
               "digital wallet and logistics modules.",
               "Implemented granular RBAC including per-seller shop roles; integrated four payment gateways "
               "(Stripe, Razorpay, MercadoPago, Telr), external shipping and two-way Odoo ERP sync.",
               "Built Elasticsearch-backed search and Redis caching; ran the platform on Octane with Horizon "
               "supervising queues.",
               "Optimised MySQL schemas and queries for high-traffic modules; moved heavy workloads into queues, "
               "background jobs and scheduled tasks.",
               "Managed cloud storage on AWS S3 and Cloudinary; contributed to architecture decisions and feature planning."]},
        {"r": "IT Sales Specialist", "c": "Global 4 Laptop Shop", "l": "Damascus, Syria", "d": "2026 – Present",
         "p": ["Advise business and individual customers on hardware, translating a stated budget and workload into "
               "a concrete specification; pre-sales consultation and post-sale technical support.",
               "Held alongside ongoing independent development work."]},
        {"r": "IT Support Specialist", "c": "BeeTronix", "l": "Aleppo, Syria", "d": "2020 – 2021",
         "p": ["Hardware, software and network support for company workstations; system installation, configuration "
               "and preventive maintenance.",
               "Documented recurring faults and fixes to speed up resolution across the support team."]},
    ],
    "proj": [
        ("EvoTech Core", "Laravel 12 · PHP 8.4 · Next.js 16 · React 19 · TypeScript",
         "Licensing, subscription and billing backbone for a product ecosystem. 14 domain modules, no cross-module "
         "joins, governed by a written architecture constitution and ADRs. URI-versioned /api/v1 with standard "
         "response envelopes and machine-readable errors; event-driven side effects."),
        ("Trydos Market", "Laravel 12 · DDD · Elasticsearch · Redis · Octane",
         "Large multi-vendor marketplace with POS, wallet and logistics. Granular per-seller RBAC, four payment "
         "gateways, Odoo ERP sync, heavy scheduling and comprehensive audit logging."),
        ("Al Ghadeer Restaurant ERP", "Laravel 12 · Blade · React 18 (Vite)",
         "Bilingual restaurant ERP serving a Blade admin panel and a React ordering SPA from one API. Audience-split "
         "controllers, session isolation so customer traffic cannot corrupt admin auth, queued FCM dispatch, "
         "versioned product cache with model-event busting."),
        ("Medical Inventory ERP", "Laravel 12 · Passport OAuth2 · MySQL",
         "Pharmaceutical distribution ERP — stock, sales orders, collections and accounting. Clean layered "
         "architecture with authorization policies and role-based scoping approximating light multi-tenancy."),
        ("Fawateer", "Flutter · Clean Architecture · BLoC · Drift",
         "Offline-first Arabic POS: live barcode scanning, Bluetooth thermal receipts, customer debt ledger, and a "
         "subscription gate with a 72-hour offline grace period."),
    ],
}

# ---------------------------------------------------- product / english ----
CV[("product", "en")] = {
    "title": "Technical Product Manager / Product Owner",
    "labels": {"summary": "Professional summary", "stack": "What I own",
               "exp": "Professional experience", "proj": "Selected products",
               "edu": "Education", "lang": "Languages", "present": "Present"},
    "summary":
        "Product owner with an engineering background: six years building and owning commercial software for small "
        "and mid-sized businesses, where I was the person who sat with the customer, decided what shipped, priced it, "
        "released it and answered the support call. I have run the full loop — discovery, scope, data model, "
        "licensing and pricing model, release planning, adoption and support — across pharmacy distribution, "
        "restaurant, retail POS and marketplace domains. Computer Systems Engineer, graduated with honours, so I can "
        "hold a credible technical conversation with the team and translate it into business terms for the customer.",
    "stack": [
        ("Discovery & definition", "Customer interviews on site, requirements elicitation from non-technical owners, "
                                   "problem framing, scope negotiation, writing acceptance criteria, prioritising a first release"),
        ("Commercial", "Subscription and licensing model design, trial and grace-period policy, multi-currency pricing "
                       "(USD/local), packaging and tiering, churn and bypass risk"),
        ("Delivery", "Release planning, versioned API policy, migration and rollout risk, staged deployment, "
                     "production support and incident triage, user documentation and training"),
        ("Stakeholders", "Direct owner-level relationships, remote cross-functional teamwork, pre-sales technical "
                         "consultation, requirement translation between business and engineering"),
        ("Domains", "Pharmaceutical distribution, restaurant and food service, retail POS and inventory, "
                    "multi-vendor e-commerce, field sales, accounting and debt ledgers"),
        ("Technical fluency", "Laravel / PHP, React / TypeScript, REST API design, MySQL, Flutter, "
                              "system architecture, embedded control systems"),
    ],
    "exp": [
        {"r": "Product Owner & Full-Stack Developer (Independent)", "c": "Self-employed",
         "l": "Damascus, Syria · Remote", "d": "2021 – Present",
         "p": ["Own a portfolio of nine commercial products end to end — discovery with the business owner through "
               "build, pricing, release and ongoing support. No spec is handed to me; I write it.",
               "Ran discovery directly with pharmacy owners, restaurateurs, shopkeepers and distributors, translating "
               "informal descriptions of how a business works into a data model and a first releasable scope.",
               "Designed the commercial model for each product: subscription licensing bound to device identity, "
               "trial periods, a 72-hour offline grace period balancing revenue protection against usability for "
               "customers with unreliable connectivity, and dual USD/local-currency pricing for an unstable market.",
               "Built EvoTech Core, the licensing, subscription and billing platform underneath the product line — "
               "turning per-product ad-hoc licensing into one governed system with automatic license issue and invoicing.",
               "Set release policy: versioned APIs and append-only migrations, so upgrades never risk live customer data.",
               "Handle production support and incident triage directly with customers, feeding recurring issues back "
               "into the roadmap."]},
        {"r": "Full-Stack Developer", "c": "Rammaz Software", "l": "Remote — Turkey", "d": "2021 – May 2026",
         "p": ["Contributed to architecture decisions and feature planning for commercial platforms across several "
               "business sectors, in a fully remote engineering team.",
               "Delivered Trydos Market, a large multi-vendor marketplace — seller onboarding and per-seller roles, "
               "orders, wallet, returns, invoicing, POS and logistics — coordinating scope across many bounded domains.",
               "Specified and integrated four payment gateways across different markets, external shipping, and "
               "two-way ERP synchronisation with Odoo, negotiating the constraints each third party imposed.",
               "Shipped a digital wallet product and financial/accounting modules where correctness and auditability "
               "were the requirement, not a nice-to-have."]},
        {"r": "IT Sales Specialist", "c": "Global 4 Laptop Shop", "l": "Damascus, Syria", "d": "2026 – Present",
         "p": ["Consultative technical sales: elicit what a customer actually needs from a vague brief and a fixed "
               "budget, then specify and justify a solution — the same discovery loop as product work, compressed "
               "into one conversation.",
               "Pre-sales consultation and post-sale support. Held alongside ongoing independent product work."]},
        {"r": "Laptop Sales Manager & IT Technician", "c": "iApple Shop", "l": "Aleppo, Syria", "d": "2021 – 2022",
         "p": ["Managed a product line end to end — selection, customer advice, and the service and support attached to it."]},
        {"r": "IT Support Specialist", "c": "BeeTronix", "l": "Aleppo, Syria", "d": "2020 – 2021",
         "p": ["Frontline support; documented recurring faults and fixes to reduce repeat resolution time across the team."]},
    ],
    "proj": [
        ("EvoTech Core — licensing & billing platform", "Product owner and engineer",
         "Problem: every product in the line reinvented its own licensing, and none of it was auditable. "
         "Built one control plane issuing licenses and invoices automatically on subscription activation, with a "
         "governed module structure so new products plug in without touching existing ones."),
        ("Fawateer — retail POS", "Product owner and engineer",
         "Problem: small shopkeepers needed billing that keeps working when the connection dies, and credit sales "
         "tracked per customer. Decided the subscription gate would allow a 72-hour offline grace period — protecting "
         "revenue without stranding a shop mid-sale — and shipped Arabic-first, since the users are not English speakers."),
        ("Smart Agent — pharmaceutical field sales", "Product owner and engineer",
         "Problem: medical reps take orders in pharmacies with no signal and reconcile later. Shipped a fully offline "
         "catalogue and order tool with PDF and Bluetooth-printed invoices, and Google Drive backup because rep phones "
         "get lost or replaced."),
        ("Trydos Market — multi-vendor marketplace", "Contributor, Rammaz Software",
         "Large marketplace spanning seller onboarding, orders, wallet, returns, invoicing, POS and logistics, with "
         "four payment gateways across different markets and Odoo ERP synchronisation."),
        ("Automatic call-to-prayer system", "Owner, installed in a mosque",
         "Non-software product with a real user constraint: the mosque staff are not technical, and the system had to "
         "run unattended for years. Designed for zero maintenance, with manual override so staff never lose control "
         "of their own sound system."),
    ],
}

# --------------------------------------------------- engineer / arabic -----
CV[("engineer", "ar")] = {
    "title": "مطوّر Full-Stack — Laravel · React · TypeScript",
    "labels": {"summary": "الملخّص المهني", "stack": "التقنيات الأساسية",
               "exp": "الخبرات المهنية", "proj": "منصّات مختارة",
               "edu": "التعليم", "lang": "اللغات", "present": "حتى الآن"},
    "summary":
        "مطوّر Full-Stack بستّ سنوات في بناء برمجيات تجارية تعتمد عليها شركات فعلية يومياً. "
        "Laravel و PHP 8 في الخلفية — طبقات خدمات مبنية على حدود النطاق، وواجهات REST مُصدَّرة بإصدارات، "
        "وتفويض عبر Policies، وطوابير وتخزين مؤقّت — مع React و Next.js و TypeScript في الواجهة. "
        "أخذتُ أربع منصّات وخمسة تطبيقات موبايل من أول متطلَّب حتى دعم الإنتاج، منها سوق متعدّد البائعين "
        "ومنصّة ترخيص وفوترة مبنية كمونوليث معياري من 14 وحدة. "
        "أعمل ضمن مساري Full-Stack: مسار Laravel مع React وهو قيد الإنتاج اليوم، ومسار TypeScript الكامل مع NestJS وأبنيه حالياً على نفس البنية. "
        "مهندس نظم حاسوبية بمرتبة الشرف، "
        "وأبني بالعربية أولاً (RTL) افتراضياً.",
    "stack": [
        ("الخلفية — Laravel", "PHP 8.4، Laravel 12، تصميم REST API، Domain-Driven Design، مونوليث معياري، "
                              "طبقة خدمات، Policies و RBAC، Sanctum / Passport، طوابير ومهام مجدولة، Octane و Horizon، WebSockets، RabbitMQ"),
        ("الواجهة — React و TypeScript", "TypeScript، React 18/19، Next.js 16، TanStack Query، "
                               "Tailwind v4، shadcn/ui، next-intl (RTL)، Blade، JavaScript، HTML5، CSS3"),
        ("Full-Stack JS (قيد التعلّم)", "NestJS، Node.js، Modules و Dependency Injection، DTOs "
                               "و Validation Pipes، Guards و Interceptors، TypeORM / Prisma، Jest — نفس البنية الطبقية "
                               "وحدود DTO والصلاحيات حسب الدور التي أعمل بها في Laravel، منقولة إلى بيئة Node. "
                               "واجهة React و TypeScript موجودة فعلياً في الإنتاج."),
        ("البيانات والبنية", "MySQL، SQLite، Redis، Elasticsearch، تصميم المخطّطات وتحسين الاستعلامات، "
                             "Nginx، Apache، Linux، PM2، AWS S3، Cloudinary، Git"),
        ("الممارسات", "بنية طبقية، حدود DTO، أحداث النطاق، إصدارات API، ترحيلات إضافية فقط، "
                      "Larastan (أقصى مستوى)، Pint، CI، مراجعة الكود"),
        ("الموبايل", "Flutter، Dart، BLoC، Provider، get_it، Clean Architecture، Drift، sqflite، بنية تعمل دون اتصال"),
        ("إضافةً إلى", "Arduino، ESP32، C/C++، أنظمة تحكّم مدمجة"),
    ],
    "exp": [
        {"r": "مطوّر Full-Stack (عمل مستقل)", "c": "لحسابي الخاص", "l": "دمشق، سوريا · عن بُعد",
         "d": "2021 – حتى الآن",
         "p": ["أملك منتجات تجارية من البداية للنهاية لصيدليات ومطاعم ومحلات تجزئة وشركات توزيع — "
               "المتطلّبات، والبنية، والتنفيذ، والإصدار، ودعم الإنتاج.",
               "بناء EvoTech Core، منصّة ترخيص واشتراكات وفوترة: مونوليث معياري من 14 وحدة نطاق تُكتشف تلقائياً "
               "وتتواصل عبر العقود وأحداث النطاق فقط، بحدود رباعية الطبقات صارمة، وعقود DTO، وثلاثة جماهير مصادقة، "
               "ومفاتيح مسار UUIDv7.",
               "واجهة أمامية بـ Next.js 16 / React 19 / TypeScript — تنميط صارم، و TanStack Query، و Tailwind v4، "
               "و next-intl، بالعربية RTL افتراضياً، منشورة على نطاقات فرعية منفصلة عبر PM2 و Nginx.",
               "ضبط الجودة عبر Larastan بأقصى مستوى و Pint و CI، وتوثيق الواجهات بـ Scramble (OpenAPI 3.1).",
               "تسليم خمسة تطبيقات Flutter بطبقات بيانات تعمل دون اتصال على SQLite و Drift، وترحيلات مُصدَّرة، "
               "وطباعة حرارية بالبلوتوث، وترخيص اشتراك مرتبط بالجهاز."]},
        {"r": "مطوّر Laravel — Full-Stack", "c": "Rammaz Software", "l": "عن بُعد — تركيا",
         "d": "2021 – أيار 2026",
         "p": ["بناء وصيانة أنظمة Laravel خلفية قابلة للتوسّع وواجهات REST مُصدَّرة تخدم عملاء الويب والموبايل، "
               "ضمن فريق يعمل عن بُعد بالكامل.",
               "إطلاق Trydos Market، منصّة واسعة متعدّدة البائعين مبنية بـ Domain-Driven Design — سياقات منفصلة "
               "للطلبات والمنتجات والمحفظة والمرتجعات والبائعين والفواتير — مع نقطة بيع ومحفظة رقمية ولوجستيات.",
               "تنفيذ صلاحيات RBAC دقيقة تشمل أدواراً على مستوى متجر كل بائع، ودمج أربع بوّابات دفع "
               "(Stripe و Razorpay و MercadoPago و Telr)، وشحن خارجي، ومزامنة ثنائية مع Odoo ERP.",
               "بناء بحث على Elasticsearch وتخزين مؤقّت على Redis، وتشغيل المنصّة على Octane مع Horizon.",
               "تحسين مخطّطات واستعلامات MySQL للوحدات عالية الحِمل، ونقل الأعباء الثقيلة إلى الطوابير والمهام المجدولة.",
               "إدارة التخزين السحابي على AWS S3 و Cloudinary، والمساهمة في قرارات البنية وتخطيط الميزات."]},
        {"r": "أخصائي مبيعات تقنية", "c": "Global 4 Laptop Shop", "l": "دمشق، سوريا", "d": "2026 – حتى الآن",
         "p": ["تقديم الاستشارة للعملاء وترجمة الميزانية وطبيعة العمل إلى مواصفات محدّدة، مع دعم تقني بعد البيع.",
               "بالتوازي مع استمرار العمل التطويري المستقل."]},
        {"r": "أخصائي دعم تقني", "c": "BeeTronix", "l": "حلب، سوريا", "d": "2020 – 2021",
         "p": ["دعم عتادي وبرمجي وشبكي لمحطات العمل، وتنصيب الأنظمة وإعدادها والصيانة الوقائية.",
               "توثيق الأعطال المتكرّرة وحلولها لتسريع المعالجة على مستوى الفريق."]},
    ],
    "proj": [
        ("EvoTech Core", "Laravel 12 · PHP 8.4 · Next.js 16 · React 19 · TypeScript",
         "العمود الفقري للترخيص والاشتراكات والفوترة لمنظومة منتجات. 14 وحدة نطاق دون أي ربط مباشر بينها، "
         "يحكمها دستور معماري رسمي وقرارات ADR. مسار ‎/api/v1‎ مُصدَّر بأغلفة استجابة قياسية وأخطاء قابلة للقراءة آلياً."),
        ("Trydos Market", "Laravel 12 · DDD · Elasticsearch · Redis · Octane",
         "سوق واسع متعدّد البائعين مع نقطة بيع ومحفظة ولوجستيات. صلاحيات دقيقة لكل بائع، وأربع بوّابات دفع، "
         "ومزامنة Odoo ERP، وجدولة مكثّفة وتسجيل تدقيق شامل."),
        ("نظام ERP لمطعم الغدير", "Laravel 12 · Blade · React 18 (Vite)",
         "نظام ERP ثنائي اللغة يخدم لوحة إدارة Blade وتطبيق طلبات React من واجهة API واحدة. متحكّمات مفصولة "
         "حسب الجمهور، وعزل جلسات يمنع حركة الزبائن من إفساد جلسة الإدارة، وإشعارات FCM عبر الطابور."),
        ("نظام ERP للتوزيع الدوائي", "Laravel 12 · Passport OAuth2 · MySQL",
         "مخزون وطلبات بيع وتحصيلات ومحاسبة. بنية طبقية نظيفة مع Policies للتفويض وتحديد نطاق حسب الدور "
         "يقارب تعدّد المستأجرين الخفيف."),
        ("فواتير (Fawateer)", "Flutter · Clean Architecture · BLoC · Drift",
         "نقطة بيع عربية تعمل دون اتصال: مسح باركود مستمر، وإيصالات حرارية بالبلوتوث، ودفتر ديون للزبائن، "
         "وبوّابة اشتراك بمهلة 72 ساعة دون اتصال."),
    ],
}

# ---------------------------------------------------- product / arabic -----
CV[("product", "ar")] = {
    "title": "مدير منتج تقني / Product Owner",
    "labels": {"summary": "الملخّص المهني", "stack": "ما أتولّاه",
               "exp": "الخبرات المهنية", "proj": "منتجات مختارة",
               "edu": "التعليم", "lang": "اللغات", "present": "حتى الآن"},
    "summary":
        "مالك منتج بخلفية هندسية: ستّ سنوات في بناء وامتلاك برمجيات تجارية للشركات الصغيرة والمتوسّطة، "
        "كنتُ فيها الشخص الذي يجلس مع الزبون، ويقرّر ما الذي يصدر، ويسعّره، ويُطلقه، ويستقبل مكالمة الدعم. "
        "أدرتُ الدورة كاملة — الاستكشاف، وتحديد النطاق، ونموذج البيانات، ونموذج الترخيص والتسعير، وتخطيط الإصدار، "
        "والتبنّي والدعم — في مجالات التوزيع الدوائي والمطاعم ونقاط البيع والأسواق الإلكترونية. "
        "مهندس نظم حاسوبية بمرتبة الشرف، أستطيع إدارة نقاش تقني حقيقي مع الفريق وترجمته إلى لغة العمل مع الزبون.",
    "stack": [
        ("الاستكشاف والتحديد", "مقابلات ميدانية مع الزبائن، واستخراج المتطلّبات من أصحاب أعمال غير تقنيين، "
                               "وصياغة المشكلة، والتفاوض على النطاق، وكتابة معايير القبول، وترتيب أولويات أول إصدار"),
        ("الجانب التجاري", "تصميم نماذج الاشتراك والترخيص، وسياسات التجربة ومهلة السماح، والتسعير متعدّد العملات "
                           "(دولار/محلي)، والتغليف والفئات، ومخاطر التسرّب والالتفاف"),
        ("التسليم", "تخطيط الإصدارات، وسياسة إصدارات API، ومخاطر الترحيل والطرح، والنشر المرحلي، "
                    "ودعم الإنتاج وفرز الحوادث، وتوثيق المستخدم والتدريب"),
        ("أصحاب المصلحة", "علاقات مباشرة على مستوى المالك، وعمل ضمن فرق موزّعة عن بُعد، واستشارة تقنية قبل البيع، "
                          "وترجمة المتطلّبات بين العمل والهندسة"),
        ("المجالات", "التوزيع الدوائي، والمطاعم والأغذية، ونقاط البيع والمخزون، والتجارة متعدّدة البائعين، "
                     "والمبيعات الميدانية، والمحاسبة ودفاتر الديون"),
        ("الطلاقة التقنية", "Laravel / PHP، React / TypeScript، تصميم REST API، MySQL، Flutter، "
                            "بنية الأنظمة، أنظمة التحكّم المدمجة"),
    ],
    "exp": [
        {"r": "Product Owner ومطوّر Full-Stack (عمل مستقل)", "c": "لحسابي الخاص",
         "l": "دمشق، سوريا · عن بُعد", "d": "2021 – حتى الآن",
         "p": ["أملك محفظة من تسعة منتجات تجارية من البداية للنهاية — من الاستكشاف مع صاحب العمل حتى البناء "
               "والتسعير والإصدار والدعم المستمر. لا يُسلَّم لي دفتر مواصفات؛ أنا من يكتبه.",
               "إدارة الاستكشاف مباشرةً مع أصحاب الصيدليات والمطاعم والمحلات وشركات التوزيع، وترجمة وصف غير رسمي "
               "لطريقة عمل النشاط إلى نموذج بيانات ونطاق أول إصدار قابل للتسليم.",
               "تصميم النموذج التجاري لكل منتج: ترخيص اشتراك مرتبط بهوية الجهاز، وفترات تجريبية، ومهلة سماح 72 ساعة "
               "دون اتصال توازن بين حماية الإيراد وقابلية الاستخدام لزبائن شبكتهم غير مستقرّة، وتسعير مزدوج "
               "بالدولار والعملة المحلية لسوق متقلّب.",
               "بناء EvoTech Core، منصّة الترخيص والاشتراكات والفوترة تحت خطّ المنتجات — تحويل ترخيص مرتجل لكل منتج "
               "إلى نظام واحد محكوم يُصدر الرخص والفواتير تلقائياً.",
               "وضع سياسة الإصدار: واجهات مُصدَّرة وترحيلات إضافية فقط، بحيث لا تعرّض الترقيات بيانات العملاء الحيّة للخطر.",
               "التعامل مع دعم الإنتاج وفرز الحوادث مباشرةً مع الزبائن، وإعادة تغذية المشكلات المتكرّرة إلى خطّة المنتج."]},
        {"r": "مطوّر Full-Stack", "c": "Rammaz Software", "l": "عن بُعد — تركيا", "d": "2021 – أيار 2026",
         "p": ["المساهمة في قرارات البنية وتخطيط الميزات لمنصّات تجارية في عدّة قطاعات، ضمن فريق يعمل عن بُعد بالكامل.",
               "تسليم Trydos Market، سوق واسع متعدّد البائعين — انضمام البائعين وأدوارهم، والطلبات والمحفظة "
               "والمرتجعات والفوترة ونقطة البيع واللوجستيات — بتنسيق النطاق عبر مجالات متعدّدة.",
               "تحديد ودمج أربع بوّابات دفع في أسواق مختلفة، وشحن خارجي، ومزامنة ثنائية مع Odoo ERP، "
               "مع التفاوض على القيود التي يفرضها كل طرف ثالث.",
               "إطلاق منتج محفظة رقمية ووحدات مالية ومحاسبية كانت الدقّة وقابلية التدقيق فيها متطلّباً أساسياً."]},
        {"r": "أخصائي مبيعات تقنية", "c": "Global 4 Laptop Shop", "l": "دمشق، سوريا", "d": "2026 – حتى الآن",
         "p": ["مبيعات تقنية استشارية: استخراج ما يحتاجه الزبون فعلاً من وصف مبهم وميزانية ثابتة، ثم تحديد الحل "
               "وتبريره — نفس دورة الاستكشاف في عمل المنتج، مضغوطة في محادثة واحدة.",
               "استشارة قبل البيع ودعم بعده، بالتوازي مع استمرار العمل المستقل على المنتجات."]},
        {"r": "مدير مبيعات لابتوبات وفنّي تقني", "c": "iApple Shop", "l": "حلب، سوريا", "d": "2021 – 2022",
         "p": ["إدارة خطّ منتجات كاملاً — الاختيار، وإرشاد الزبون، والخدمة والدعم المرتبطين به."]},
        {"r": "أخصائي دعم تقني", "c": "BeeTronix", "l": "حلب، سوريا", "d": "2020 – 2021",
         "p": ["دعم مباشر للمستخدمين، وتوثيق الأعطال المتكرّرة وحلولها لتقليل زمن المعالجة المتكرّر على مستوى الفريق."]},
    ],
    "proj": [
        ("EvoTech Core — منصّة ترخيص وفوترة", "مالك المنتج والمهندس",
         "المشكلة: كل منتج في الخطّ كان يعيد اختراع ترخيصه الخاص، ولا شيء منه قابل للتدقيق. بنيتُ منصّة واحدة "
         "تُصدر الرخص والفواتير تلقائياً عند تفعيل الاشتراك، ببنية وحدات محكومة تسمح بإضافة منتجات جديدة "
         "دون المساس بالقائم."),
        ("فواتير — نقطة بيع للتجزئة", "مالك المنتج والمهندس",
         "المشكلة: أصحاب المحال الصغيرة يحتاجون فوترة تستمر بالعمل عند انقطاع الشبكة، وتتبّعاً للبيع بالدَّين لكل زبون. "
         "قرّرتُ أن تسمح بوّابة الاشتراك بمهلة 72 ساعة دون اتصال — تحمي الإيراد دون أن تُعطّل محلاً في منتصف عملية بيع — "
         "وأُطلق بالعربية أولاً، لأن المستخدمين لا يتحدّثون الإنكليزية."),
        ("المندوب الذكي — مبيعات دوائية ميدانية", "مالك المنتج والمهندس",
         "المشكلة: مندوبو الأدوية يأخذون الطلبات في صيدليات دون تغطية ويسوّونها لاحقاً. أطلقتُ كتالوغاً وأداة طلبات "
         "تعمل دون اتصال بالكامل، بفواتير PDF وطباعة بلوتوث، ونسخ احتياطي على Google Drive لأن هواتف المندوبين "
         "تُفقد أو تُستبدل."),
        ("Trydos Market — سوق متعدّد البائعين", "مساهم، Rammaz Software",
         "سوق واسع يغطّي انضمام البائعين والطلبات والمحفظة والمرتجعات والفوترة ونقطة البيع واللوجستيات، "
         "مع أربع بوّابات دفع في أسواق مختلفة ومزامنة مع Odoo ERP."),
        ("نظام الأذان الآلي", "مالك المنتج، منفَّذ في مسجد",
         "منتج غير برمجي بقيد مستخدم حقيقي: القائمون على المسجد ليسوا تقنيين، والنظام يجب أن يعمل سنوات دون تدخّل. "
         "صُمّم لصفر صيانة، مع تجاوز يدوي كي لا يفقد الطاقم التحكّم بنظام الصوت الخاص بهم."),
    ],
}

EDU = {
    "en": {"degree": "Bachelor of Computer Systems Engineering (Automation & Control)",
           "school": "University of Aleppo", "date": "Graduated 2020",
           "note": "GPA 89.7 / 100 — Excellent with Honours",
           "langs": [("Arabic", "Native"), ("English", "Professional working proficiency")]},
    "ar": {"degree": "إجازة في هندسة النظم الحاسوبية (أتمتة وتحكّم)",
           "school": "جامعة حلب", "date": "التخرّج 2020",
           "note": "المعدّل 89.7 / 100 — امتياز مع مرتبة الشرف",
           "langs": [("العربية", "اللغة الأم"), ("الإنكليزية", "إجادة عملية احترافية")]},
}

# ================================================================== PDF =====

CSS = """
@page { size: A4; margin: 13mm 14mm 12mm; }
* { margin:0; padding:0; box-sizing:border-box; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family:"Inter","Segoe UI",sans-serif; font-size:9.1pt; line-height:1.5; color:#1a1a1a; }
body.rtl { direction:rtl; font-family:"Segoe UI","Tahoma",sans-serif; }
a { color:#0d5c73; text-decoration:none; }

.name { font-size:21pt; font-weight:700; letter-spacing:-.02em; color:#0f1720; }
.role { font-size:10.6pt; color:#0d5c73; font-weight:600; margin-top:3pt; }
.contact { margin-top:6pt; font-size:8.5pt; color:#4a5568; }
.contact span { white-space:nowrap; }
.contact .sep { color:#b8c2cf; padding:0 5pt; }
.links { margin-top:2.5pt; font-size:8.5pt; }
hr.rule { border:none; border-top:1.6pt solid #0d5c73; margin:9pt 0 0; }

h2 { font-size:8.2pt; font-weight:700; letter-spacing:.15em; text-transform:uppercase;
     color:#0d5c73; margin:11pt 0 5pt; border-bottom:.6pt solid #d6dee8; padding-bottom:2.5pt; }
body.rtl h2 { letter-spacing:0; }

.summary { text-align:justify; color:#2d3748; }

.srow { display:flex; gap:7pt; margin-bottom:3.2pt; page-break-inside:avoid; }
.srow .k { flex:0 0 30mm; font-weight:700; color:#0f1720; font-size:8.6pt; }
.srow .v { flex:1; color:#3a4658; font-size:8.6pt; }

.job { margin-bottom:7.5pt; page-break-inside:avoid; }
.job .top { display:flex; justify-content:space-between; align-items:baseline; gap:8pt; }
.job .r { font-weight:700; font-size:9.9pt; color:#0f1720; }
.job .d { font-size:8.3pt; color:#5a6675; white-space:nowrap; font-weight:600; }
.job .co { font-size:8.9pt; color:#0d5c73; font-weight:600; margin-top:1pt; }
.job .co .loc { color:#6b7785; font-weight:400; }
ul { list-style:none; margin-top:3.5pt; }
li { position:relative; padding-left:8.5pt; margin-bottom:2.6pt; color:#2d3748; text-align:justify; }
body.rtl li { padding-left:0; padding-right:8.5pt; }
li::before { content:""; position:absolute; left:1.5pt; top:4.2pt; width:2.6pt; height:2.6pt;
             border-radius:50%; background:#0d5c73; }
body.rtl li::before { left:auto; right:1.5pt; }

.proj { margin-bottom:5.5pt; page-break-inside:avoid; }
.proj .t { font-weight:700; font-size:9.2pt; color:#0f1720; }
.proj .s { font-size:8pt; color:#0d5c73; font-weight:600; }
.proj .b { font-size:8.5pt; color:#3a4658; margin-top:1pt; text-align:justify; }

.two { display:flex; gap:9mm; }
.two > div { flex:1; }
.edu .d { font-weight:700; font-size:9pt; color:#0f1720; }
.edu .m { font-size:8.5pt; color:#4a5568; }
.edu .n { font-size:8.5pt; color:#0d5c73; font-weight:600; }
.lrow { display:flex; justify-content:space-between; font-size:8.6pt; margin-bottom:2pt; }
.lrow .k { font-weight:600; color:#0f1720; }
.lrow .v { color:#4a5568; }
"""

def build_html(variant, lang):
    c, e = CV[(variant, lang)], EDU[lang]
    L, x = c["labels"], html.escape
    rtl = " rtl" if lang == "ar" else ""

    stack = "".join('<div class="srow"><div class="k">%s</div><div class="v">%s</div></div>'
                    % (x(k), x(v)) for k, v in c["stack"])

    jobs = "".join(
        '<div class="job"><div class="top"><span class="r">%s</span><span class="d">%s</span></div>'
        '<div class="co">%s<span class="loc"> · %s</span></div><ul>%s</ul></div>'
        % (x(j["r"]), x(j["d"]), x(j["c"]), x(j["l"]),
           "".join("<li>%s</li>" % x(p) for p in j["p"]))
        for j in c["exp"])

    projs = "".join('<div class="proj"><div class="t">%s <span class="s">— %s</span></div>'
                    '<div class="b">%s</div></div>' % (x(t), x(s), x(b))
                    for t, s, b in c["proj"])

    langs = "".join('<div class="lrow"><span class="k">%s</span><span class="v">%s</span></div>'
                    % (x(k), x(v)) for k, v in e["langs"])

    return """<!doctype html><html lang="%s" dir="%s"><head><meta charset="utf-8">
<title>%s</title><style>%s</style><style>%s</style></head><body class="%s">
<div class="name">%s</div>
<div class="role">%s</div>
<div class="contact"><span>%s</span><span class="sep">|</span><span dir="ltr">%s</span><span class="sep">|</span><span dir="ltr">%s</span></div>
<div class="contact links" dir="ltr"><span>%s</span><span class="sep">|</span><span>%s</span><span class="sep">|</span><span>%s</span></div>
<hr class="rule">
<h2>%s</h2><div class="summary">%s</div>
<h2>%s</h2>%s
<h2>%s</h2>%s
<h2>%s</h2>%s
<div class="two"><div><h2>%s</h2><div class="edu"><div class="d">%s</div>
<div class="m">%s · %s</div><div class="n">%s</div></div></div>
<div><h2>%s</h2>%s</div></div>
</body></html>""" % (
        lang, "rtl" if lang == "ar" else "ltr",
        x(PROF["name"][lang]), FONTS, CSS, rtl.strip(),
        x(PROF["name"][lang]), x(c["title"]),
        x(PROF["location"][lang]), x(PROF["phone"]), x(PROF["email"]),
        SITE, GITHUB, LINKEDIN,
        x(L["summary"]), x(c["summary"]),
        x(L["stack"]), stack,
        x(L["exp"]), jobs,
        x(L["proj"]), projs,
        x(L["edu"]), x(e["degree"]), x(e["school"]), x(e["date"]), x(e["note"]),
        x(L["lang"]), langs)

FONTS = io.open(os.path.join(SP, "fonts.css"), encoding="utf-8").read()

CHROME = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
          r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
          "google-chrome", "chromium"]

def to_pdf(doc, path):
    hp = os.path.join(tempfile.gettempdir(), "cv_%s.html" % os.path.basename(path))
    io.open(hp, "w", encoding="utf-8").write(doc)
    exe = next((c for c in CHROME if os.path.exists(c) or not c.startswith("C:")), None)
    subprocess.run([exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    "--virtual-time-budget=20000", "--print-to-pdf=" + path,
                    "file:///" + hp.replace("\\", "/")], check=True, capture_output=True)

# ================================================================= DOCX =====

def to_docx(variant, lang, path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    c, e = CV[(variant, lang)], EDU[lang]
    L, rtl = c["labels"], lang == "ar"
    ACCENT, DARK, MUT = RGBColor(0x0D, 0x5C, 0x73), RGBColor(0x0F, 0x17, 0x20), RGBColor(0x4A, 0x55, 0x68)
    face = "Segoe UI" if rtl else "Calibri"

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.3)
        s.left_margin = s.right_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = face
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:cs"), face)

    def para(space_after=3, align=None):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(space_after)
        pf.line_spacing = 1.12
        if rtl:
            bidi = OxmlElement("w:bidi"); p._p.get_or_add_pPr().append(bidi)
            pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if align is not None:
            pf.alignment = align
        return p

    def run(p, text, size=9.5, bold=False, color=None, caps=False):
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
        r.font.name = face
        r._element.rPr.rFonts.set(qn("w:cs"), face)
        if color is not None:
            r.font.color.rgb = color
        if caps:
            r.font.all_caps = True
        return r

    def heading(text):
        p = para(space_after=2)
        p.paragraph_format.space_before = Pt(9)
        run(p, text, 8.5, True, ACCENT, caps=not rtl)
        pr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
        bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "4")
        bt.set(qn("w:space"), "2"); bt.set(qn("w:color"), "D6DEE8")
        bd.append(bt); pr.append(bd)

    # header
    p = para(space_after=1); run(p, PROF["name"][lang], 20, True, DARK)
    p = para(space_after=4); run(p, c["title"], 10.5, True, ACCENT)
    p = para(space_after=1)
    run(p, "%s  |  %s  |  %s" % (PROF["location"][lang], PROF["phone"], PROF["email"]), 8.5, color=MUT)
    p = para(space_after=6)
    run(p, "%s  |  %s  |  %s" % (SITE, GITHUB, LINKEDIN), 8.5, color=ACCENT)

    heading(L["summary"])
    p = para(space_after=4); run(p, c["summary"], 9.3, color=MUT)

    heading(L["stack"])
    for k, v in c["stack"]:
        p = para(space_after=2)
        run(p, k + ": ", 9, True, DARK)
        run(p, v, 9, color=MUT)

    heading(L["exp"])
    for j in c["exp"]:
        p = para(space_after=0)
        run(p, j["r"], 10, True, DARK)
        run(p, "   " + j["d"], 8.5, True, MUT)
        p = para(space_after=2)
        run(p, j["c"], 9, True, ACCENT); run(p, " · " + j["l"], 9, color=MUT)
        for b in j["p"]:
            bp = para(space_after=2)
            bp.paragraph_format.left_indent = Cm(0 if rtl else 0.45)
            bp.paragraph_format.right_indent = Cm(0.45 if rtl else 0)
            run(bp, "• " + b, 9.2, color=MUT)
        para(space_after=3)

    heading(L["proj"])
    for t, s, b in c["proj"]:
        p = para(space_after=1)
        run(p, t, 9.4, True, DARK); run(p, " — " + s, 8.6, True, ACCENT)
        p = para(space_after=4); run(p, b, 9, color=MUT)

    heading(L["edu"])
    p = para(space_after=1); run(p, e["degree"], 9.4, True, DARK)
    p = para(space_after=1); run(p, "%s · %s" % (e["school"], e["date"]), 9, color=MUT)
    p = para(space_after=5); run(p, e["note"], 9, True, ACCENT)

    heading(L["lang"])
    for k, v in e["langs"]:
        p = para(space_after=2)
        run(p, k + ": ", 9, True, DARK); run(p, v, 9, color=MUT)

    doc.save(path)

# ================================================================== main ====

NAMES = {("engineer", "en"): "Mohamad-Hasan-CV-FullStack-EN",
         ("engineer", "ar"): "Mohamad-Hasan-CV-FullStack-AR",
         ("product",  "en"): "Mohamad-Hasan-CV-Product-EN",
         ("product",  "ar"): "Mohamad-Hasan-CV-Product-AR"}

def main():
    args = sys.argv[1:]
    variants = [a for a in args if a in ("engineer", "product")] or ["engineer", "product"]
    langs    = [a for a in args if a in ("en", "ar")] or ["en", "ar"]
    if not os.path.isdir(OUT):
        os.makedirs(OUT)
    for v in variants:
        for l in langs:
            base = os.path.join(OUT, NAMES[(v, l)])
            to_pdf(build_html(v, l), base + ".pdf")
            to_docx(v, l, base + ".docx")
            print("%-38s pdf %4d KB   docx %3d KB" % (
                NAMES[(v, l)], os.path.getsize(base + ".pdf") / 1024,
                os.path.getsize(base + ".docx") / 1024))

if __name__ == "__main__":
    main()
